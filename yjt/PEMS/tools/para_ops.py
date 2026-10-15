# -*- coding: utf-8 -*-
"""
para_ops —— 段落级「重复段检测 + 版式一致性」通用原语
=====================================================
承接（上一版文字）与槽位（本轮按权威源重写）是两套独立写入轨道，二者会各自落到一个段落上，
产出「同一句话两个版本、且分别落在版式不同的段落」。本模块只提供可复用的原子操作，
不含任何项目数据；项目差异通过参数与 class_of 回调传入。

依赖：python-docx
"""
import copy
import difflib
import re

from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IND_ATTRS = ('left', 'leftChars', 'right', 'rightChars', 'firstLine',
             'firstLineChars', 'hanging', 'hangingChars')
SPC_ATTRS = ('before', 'after', 'line', 'lineRule')
TXBX = qn('w:txbxContent')
AUTHORED_KINDS = ('slot', 'extra', 'carry')


def nws(s):
    return re.sub(r'\s+', '', s or '')


# --------------------------------------------------------------------------
# 重复段
# --------------------------------------------------------------------------
def dup_kind(a, b):
    """两段是否「同一句话的两个版本」，返回类型标签或 None。四种形态缺一不可：

      exact      同一段被写进两个位置（承接把上一版一段话落到两个模板空段）
      near       同义改写，只差标点或一个连接词
      same-lead  开头若干字完全相同、后面条目列表不同 —— 整体相似度会被稀释到 0.6 以下，
                 只看 ratio 必漏，必须先比开头
      contained  一段是另一段的子串（上一版把两句合成一句，本轮又拆开写）
    """
    a, b = nws(a), nws(b)
    if not a or not b or min(len(a), len(b)) < 12:
        return None
    r = difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()
    if r >= 0.999:
        return 'exact'
    if r >= 0.85:
        return 'near'
    if a[:12] == b[:12] and r >= 0.50:
        return 'same-lead'
    if min(len(a), len(b)) >= 25 and (a in b or b in a):
        return 'contained'
    return None


def style_has_numbering(p):
    """段落自己没有 w:numPr，但样式链（如 Figure）自带编号时同样属"编号段"。
    只看段落级 pPr 会把这类列表项当正文，套上正文首行缩进（pitfalls §34）。"""
    st, seen = p.style, set()
    while st is not None and id(st) not in seen:
        seen.add(id(st))
        el = st.element.find(qn('w:pPr') + '/' + qn('w:numPr'))
        if el is not None:
            return True
        st = st.base_style
    return False


def is_structural_paragraph(p):
    """标题 / 目录 / 图题 / 含域 / 含图形 / 文本框 / 编号段（含样式链继承）：
    不参与正文重复与版式比对"""
    el = p._p
    if el.find('.//' + qn('w:fldChar')) is not None \
            or el.find('.//' + qn('w:instrText')) is not None:
        return True
    if el.findall('.//' + qn('w:pict')) or el.findall('.//' + qn('w:drawing')):
        return True
    if _in_textbox(el):
        return True
    st = (p.style.name or '').lower() if p.style else ''
    if st.startswith('heading') or 'toc' in st or st in ('title', 'caption'):
        return True
    pPr = el.find(qn('w:pPr'))
    if (pPr is not None and pPr.find(qn('w:numPr')) is not None) or style_has_numbering(p):
        return True
    t = nws(''.join(n.text or '' for n in el.iter(qn('w:t'))))
    if re.match(r'^(图|表)\s*\d', t):
        return True
    runs = [r for r in p.runs if r.text.strip()]
    if runs and len(t) < 40 and all(r.bold for r in runs):
        return True
    jc = pPr.find(qn('w:jc')) if pPr is not None else None
    if jc is not None and jc.get(qn('w:val')) == 'center' and len(t) < 60:
        return True
    return False


def find_duplicate_paragraphs(doc, allow=()):
    """正文重复段清单。allow=[(子串A, 子串B)] 声明模板自带的同族并列条目（如两种"××验证"）。"""
    items = [(i, p) for i, p in enumerate(doc.paragraphs)
             if nws(p.text) and not is_structural_paragraph(p)]
    out = []
    for a in range(len(items)):
        for b in range(a + 1, len(items)):
            na, nb = nws(items[a][1].text), nws(items[b][1].text)
            k = dup_kind(na, nb)
            if not k:
                continue
            if any((x in na and y in nb) or (y in na and x in nb) or (x in na and x in nb)
                   for x, y in allow):
                continue
            out.append((k, items[a][0], items[b][0], na[:40], nb[:40]))
    return out


def dedupe_paragraphs(doc, class_of, allow=(), log=None):
    """删除重复段，只留一条。优先级：本轮正文 > 模板原样 > 上一版承接；同级留靠前的。

    两段都是模板原样时**不自动删**：那是模板自带的写法（实测两条并列的"××验证"条目
    开头相同），删掉就是吃正文，只报给闸门由人判断。
    胜者若没带版式而败者带了，先把败者的版式接过来（留文字与留版式分别取最合适的来源）。
    """
    rank = {'slot': 3, 'extra': 3, 'base': 2, 'carry': 1}
    items = [(i, p) for i, p in enumerate(doc.paragraphs)
             if nws(p.text) and not is_structural_paragraph(p)]
    dead, dropped = set(), []
    for a in range(len(items)):
        if items[a][1]._p in dead:
            continue
        for b in range(a + 1, len(items)):
            pb = items[b][1]
            if pb._p in dead:
                continue
            na, nb = nws(items[a][1].text), nws(pb.text)
            k = dup_kind(na, nb)
            if not k:
                continue
            if any((x in na and y in nb) or (y in na and x in nb) or (x in na and x in nb)
                   for x, y in allow):
                continue
            ia, pa = items[a]
            ra, rb = rank.get(class_of(pa), 2), rank.get(class_of(pb), 2)
            if ra == rb == rank['base']:
                if log:
                    log('  [重·待判] #%d↔#%d %s（两段都是模板原样，需人工判断）' % (ia, items[b][0], k))
                continue
            loser, lidx, w, win_idx = (pb, items[b][0], pa, ia) if ra >= rb else (pa, ia, pb, items[b][0])
            if class_of(w) in AUTHORED_KINDS and _ppr_direct(w) is None and _ppr_direct(loser) is not None:
                set_para_format_like(w, loser, doc)
            dead.add(loser._p)
            dropped.append((win_idx, lidx, k))
            if loser is pa:
                break
    for el in dead:
        el.getparent().remove(el)
    for w, l, k in dropped:
        if log:
            log('  [重] %s 留#%d 删#%d' % (k, w, l))
    return len(dead)


# --------------------------------------------------------------------------
# 版式一致性
# --------------------------------------------------------------------------
def _in_textbox(el):
    a = el.getparent()
    while a is not None:
        if a.tag == TXBX:
            return True
        a = a.getparent()
    return False


def _ppr_direct(p):
    """取段落自己的 w:pPr。兼容两种入参：Paragraph 对象，或已经取出来的 CT_P 元素
    （`is_struct_carrier` 内部传的是 `p._p`，只认 Paragraph 会 AttributeError）。"""
    el = getattr(p, '_p', p)
    return el.find(qn('w:pPr')) if el is not None else None


def has_direct_format(p):
    """段落是否「自己写了」缩进/行距/对齐。空 <w:pPr/> 不算：那是模板残留的空壳，
    把它当成"刻意版式"会让整片缺缩进的段落互相充当锚点，漂移就检不出来了。"""
    pPr = _ppr_direct(p)
    return pPr is not None and any(pPr.find(qn(t)) is not None
                                   for t in ('w:ind', 'w:spacing', 'w:jc'))


def _merge_ppr(acc, pPr):
    if pPr is None:
        return
    for tag in ('w:ind', 'w:spacing', 'w:jc'):
        e = pPr.find(qn(tag))
        if e is None:
            continue
        if tag == 'w:jc':
            acc['jc'] = e.get(qn('w:val'))
            continue
        for k, v in e.attrib.items():
            acc[tag[2:] + ':' + k.split('}')[-1]] = v


def effective_para_format(doc, p):
    """段落的生效版式：docDefaults → 样式链（basedOn 逐级）→ 段落直接 pPr。

    只读段落 pPr 会把"样式里已设好缩进、段落本身不写 pPr"误判成缺格式；
    Word 的观感由继承链决定，比对必须按生效值，否则既冤枉模板又漏掉真漂移。
    """
    raw = {}
    dd = doc.styles.element.find(qn('w:docDefaults'))
    if dd is not None:
        pdd = dd.find(qn('w:pPrDefault'))
        if pdd is not None:
            _merge_ppr(raw, pdd.find(qn('w:pPr')))
    chain, st, seen = [], p.style, set()
    while st is not None and id(st) not in seen:
        seen.add(id(st))
        chain.append(st.element)
        st = st.base_style
    for el in reversed(chain):
        _merge_ppr(raw, el.find(qn('w:pPr')))
    _merge_ppr(raw, _ppr_direct(p))
    return raw


def para_format_sig(doc, p):
    """可比较的版式指纹：同义写法归一（0/c0/t0、单倍行距、未写 jc 与 both 视为相同）。"""
    raw = effective_para_format(doc, p)
    fc, ft = raw.get('ind:firstLineChars'), raw.get('ind:firstLine')
    if fc and fc != '0':
        fl = 'c' + fc
    elif ft and ft != '0':
        fl = 't' + ft
    else:
        fl = ''
    left = raw.get('ind:leftChars') or raw.get('ind:left') or ''
    left = '' if left == '0' else left
    line, rule = raw.get('spacing:line') or '', raw.get('spacing:lineRule') or ''
    if line in ('0', '240') and rule in ('', 'auto'):
        line = rule = ''
    return (fl, left, line, rule, raw.get('jc') or 'both')


def set_para_format_like(p, donor, doc):
    """让 p 的生效版式与 donor 一致：对齐 pStyle（双向）+ 对齐 donor 直接写的缩进/行距/对齐。

    三条实测必要性：
      1) 只抄 donor 的直接 pPr 不够 —— 两段都无 pPr 时差异全来自各自样式，抄不到东西；
      2) pStyle 必须双向：donor 没写（用文档默认样式）而本段写了，也要摘掉；
      3) 把 donor 的「生效值」全量写成直接属性也不行 —— 正文会冒出一堆模板没有的硬格式，
         用户在 Word 里改样式跟不动。
    """
    changed = False
    d_ppr = _ppr_direct(donor)
    t_ppr = _ppr_direct(p)
    d_style = d_ppr.find(qn('w:pStyle')) if d_ppr is not None else None
    want_style = d_style.get(qn('w:val')) if d_style is not None else None
    olds = t_ppr.findall(qn('w:pStyle')) if t_ppr is not None else []
    cur = olds[0].get(qn('w:val')) if len(olds) == 1 else None
    if cur != want_style:
        t_ppr = _ensure_ppr(p, t_ppr)
        for old in olds:
            t_ppr.remove(old)
        if want_style:
            t_ppr.insert(0, copy.deepcopy(d_style))
        changed = True
    for tag, allowed in (('w:ind', IND_ATTRS), ('w:spacing', SPC_ATTRS)):
        de = d_ppr.find(qn(tag)) if d_ppr is not None else None
        want = {} if de is None else {k.split('}')[-1]: v for k, v in de.attrib.items()
                                      if k.split('}')[-1] in allowed}
        olds = t_ppr.findall(qn(tag)) if t_ppr is not None else []
        cur = (_attrs(olds[0]) if len(olds) == 1 else None)
        tgt = tuple(sorted(want.items(), key=lambda x: allowed.index(x[0])))
        if cur == tgt or (not want and not olds):
            continue
        t_ppr = _ensure_ppr(p, t_ppr)
        for old in olds:
            t_ppr.remove(old)
        if want:
            e = OxmlElement(tag)
            for k, v in tgt:
                e.set(qn('w:' + k), str(v))
            t_ppr.append(e)
        changed = True
    dj = d_ppr.find(qn('w:jc')) if d_ppr is not None else None
    want_jc = dj.get(qn('w:val')) if dj is not None else None
    olds = t_ppr.findall(qn('w:jc')) if t_ppr is not None else []
    cur = olds[0].get(qn('w:val')) if len(olds) == 1 else None
    if cur != want_jc:
        t_ppr = _ensure_ppr(p, t_ppr)
        for old in olds:
            t_ppr.remove(old)
        if want_jc:
            e = OxmlElement('w:jc')
            e.set(qn('w:val'), want_jc)
            t_ppr.append(e)
        changed = True
    return changed


def _ensure_ppr(p, t_ppr):
    if t_ppr is not None:
        return t_ppr
    t_ppr = OxmlElement('w:pPr')
    p._p.insert(0, t_ppr)
    return t_ppr


def _attrs(el):
    return tuple(sorted((k.split('}')[-1], v) for k, v in el.attrib.items()))


def body_paragraphs(doc, min_chars=25):
    return [p for p in doc.paragraphs
            if len(nws(''.join(n.text or '' for n in p._p.iter(qn('w:t'))))) >= min_chars
            and not is_structural_paragraph(p)]


def para_format_pairs(longs, class_of):
    """给「本轮新写且自己没带版式」的正文段配版式锚点。

    锚点 = 模板原样段（base）或任何直接写了 pPr 的段（那是模板刻意给该位置的版式）。
    取法单向就近（先往前再往后），锚点自身永不改写，所以一轮就定。
    不要用「就近互指」：整片都是本轮新写时 A 抄 B、B 抄 A，几轮都定不下来；
    也不要用「±N 多数派」：会把模板自己写了缩进的段拉平成样式默认值，等于新造不一致。
    """
    hard = [n for n, q in enumerate(longs)
            if class_of(q) == 'base' or has_direct_format(q)]
    out = []
    for n, q in enumerate(longs):
        if class_of(q) not in AUTHORED_KINDS or has_direct_format(q) or n in hard:
            continue
        back = [m for m in hard if m < n]
        fwd = [m for m in hard if m > n]
        a = back[-1] if back else (fwd[0] if fwd else None)
        if a is not None:
            out.append((n, a))
    return out


def para_format_issues(doc, longs, class_of):
    """本轮正文段与锚点生效版式不一致的清单（生成器据此修，闸门据此判）"""
    out = []
    for n, a in para_format_pairs(longs, class_of):
        own, ref = para_format_sig(doc, longs[n]), para_format_sig(doc, longs[a])
        if own != ref:
            out.append((n, a, ref, own))
    return out


def align_para_formats(doc, class_of, log=None):
    """把本轮新写的正文段补齐到锚点的生效版式；返回改动数。

    class_of(p) → 'slot'/'extra'/'carry'/'base'：生成器必须记录每段文字的来历，
    否则分不清"本轮写坏了"与"模板本来就这样"，只能整片重排（越界且危险）。
    """
    longs = body_paragraphs(doc)
    fixed = 0
    for n, a, ref, _own in para_format_issues(doc, longs, class_of):
        p, anchor = longs[n], longs[a]
        if set_para_format_like(p, anchor, doc):
            fixed += 1
            if log:
                log('  [版式] 正文段#%d ← 锚点#%d %s' % (n, a, ref))
            if para_format_sig(doc, p) != ref and log:
                log('       ‼ 补齐后仍与锚点不同（样式链差异），需人工确认')
    return fixed


# --------------------------------------------------------------------------
# 分页与导航结构（章前分页、标题树、空标题段）
# --------------------------------------------------------------------------
def outline_level(p):
    """段落大纲级别（直接 pPr 或样式链）。'9' = 正文级（不进导航窗格/目录）。"""
    pPr = _ppr_direct(p)
    if pPr is not None:
        e = pPr.find(qn('w:outlineLvl'))
        if e is not None:
            return e.get(qn('w:val'))
    st, seen = p.style, set()
    while st is not None and id(st) not in seen:
        seen.add(id(st))
        e = st.element.find(qn('w:pPr') + '/' + qn('w:outlineLvl'))
        if e is not None:
            return e.get(qn('w:val'))
        st = st.base_style
    return None


def set_outline_level(p, lvl):
    """只写 outlineLvl，不改 pStyle —— 套 Heading 样式会连带改字号，违反标题基线。"""
    pPr = _ensure_ppr(p, _ppr_direct(p))
    for old in pPr.findall(qn('w:outlineLvl')):
        pPr.remove(old)
    e = OxmlElement('w:outlineLvl')
    e.set(qn('w:val'), str(lvl))
    pPr.append(e)


def set_page_break_before(p, val='1'):
    """章前分页用 pageBreakBefore，不要插 w:br type=page（会新增段落、可被误删）。"""
    pPr = _ensure_ppr(p, _ppr_direct(p))
    for old in pPr.findall(qn('w:pageBreakBefore')):
        pPr.remove(old)
    e = OxmlElement('w:pageBreakBefore')
    e.set(qn('w:val'), val)
    pPr.insert(0 if pPr.find(qn('w:pStyle')) is None else 1, e)


def has_page_break_before(p):
    pPr = _ppr_direct(p)
    e = pPr.find(qn('w:pageBreakBefore')) if pPr is not None else None
    return e is not None and e.get(qn('w:val')) not in ('0', 'false')


def is_struct_carrier(p):
    """删了会破坏结构的段落：图形载体、分节符载体、含域代码、夹在两个表格之间的隔离段。"""
    el = p._p
    if el.findall('.//' + qn('w:drawing')) or el.findall('.//' + qn('w:pict')):
        return True
    pPr = _ppr_direct(el)
    if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
        return True
    if el.find('.//' + qn('w:fldChar')) is not None or el.find('.//' + qn('w:instrText')) is not None:
        return True
    prv, nxt = el.getprevious(), el.getnext()
    return bool(prv is not None and prv.tag == qn('w:tbl')
                and nxt is not None and nxt.tag == qn('w:tbl'))


def tidy_navigation(doc, chapter_heading, log=None):
    """删空标题段与章前顶页空段；结构载体不能删的，降级为 outlineLvl=9。

    必须按最新段落表反复重扫：删一段后序号全变，一次性遍历会漏（实测连续空段只删一半）。
    """
    body = doc.element.body
    removed = demoted = 0
    for _round in range(6):
        ps = doc.paragraphs
        last = ps[-1]._p if ps else None
        did = 0
        for i, p in enumerate(ps):
            el = p._p
            if el is last or el.getparent() is not body or nws(p.text):
                continue
            lvl = outline_level(p)
            st = (p.style.name or '').lower() if p.style else ''
            heading_like = lvl not in (None, '9') or st.startswith('heading') or st.startswith('title')
            carrier = is_struct_carrier(p)
            if heading_like and carrier:
                set_outline_level(p, 9)
                demoted += 1
                continue
            if carrier:
                continue
            ahead = False
            if heading_like:
                ahead = True
            else:
                j = i + 1
                while j < len(ps) and not nws(ps[j].text):
                    j += 1
                ahead = j < len(ps) and chapter_heading(ps[j])
            if ahead:
                body.remove(el)
                removed += 1
                did += 1
        if not did:
            break
    if log:
        log('  [导航] 删除空标题段/顶页空段 %d 个，降级结构载体 %d 个' % (removed, demoted))
    return removed, demoted
