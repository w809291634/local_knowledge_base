# -*- coding: utf-8 -*-
"""
docx_ops —— Word 文档「模板就地替换」通用原语库
================================================
本模块只提供**可复用的原子操作**，不含任何项目数据。
项目差异全部通过参数 / 配置传入，从而保证「相同代码只写一份」。

依赖：python-docx
"""
import copy
import difflib
import hashlib
import re
import os
import struct

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.run import Run


# --------------------------------------------------------------------------
# 只读工具
# --------------------------------------------------------------------------
def paragraph_text(p):
    """取段落纯文本（含所有 run，忽略图内文字）"""
    return ''.join(t.text or '' for t in p._p.iter(qn('w:t')))


def has_field_code(p):
    """段落是否含域代码（目录 PAGEREF 等）；此类段落不可做 run 合并"""
    return bool(p._p.findall('.//' + qn('w:fldChar')) or
                p._p.findall('.//' + qn('w:instrText')))


def has_graphic(p):
    """段落是否含图形（DrawingML 或 VML）"""
    for tag in ('w:drawing', 'w:pict'):
        if p._p.findall('.//' + qn(tag)):
            return True
    return False


def has_section_break(p):
    """段落是否携带分节符（pPr/sectPr）"""
    pPr = p._p.find(qn('w:pPr'))
    return pPr is not None and pPr.find(qn('w:sectPr')) is not None


def find_cover_table(body, first_cell_text='文件编号'):
    """按「首行首格文本」定位封面信息表（浮动表）"""
    for tbl in body.findall(qn('w:tbl')):
        trs = tbl.findall(qn('w:tr'))
        if not trs:
            continue
        tcs = trs[0].findall(qn('w:tc'))
        if not tcs:
            continue
        cell0 = ''.join(t.text or '' for t in tcs[0].iter(qn('w:t'))).strip()
        if cell0 == first_cell_text:
            return tbl
    return None


def describe_blocks(doc):
    """打印块序列（段落/表格 + 编号/分节/图形标记），用于动手前摸清模板"""
    out = []
    for i, el in enumerate(doc.element.body, 1):
        if el.tag == qn('w:p'):
            pPr = el.find(qn('w:pPr'))
            flag = ''
            if pPr is not None:
                if pPr.find(qn('w:numPr')) is not None:
                    flag += 'NUM'
                if pPr.find(qn('w:sectPr')) is not None:
                    flag += 'SECT'
            if el.findall('.//' + qn('w:drawing')) or el.findall('.//' + qn('w:pict')):
                flag += 'IMG'
            txt = ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()
            out.append('%3d P %-10s %s' % (i, flag, txt[:70]))
        elif el.tag == qn('w:tbl'):
            out.append('%3d TBL rows=%d' % (i, len(el.findall(qn('w:tr')))))
    return '\n'.join(out)


# --------------------------------------------------------------------------
# 段落替换
# --------------------------------------------------------------------------
def set_paragraph(p, text, label=None):
    """
    替换整段文字，保留原格式。
    - label=None ：整段替换为 text（沿用首个 run 的 rPr）。
    - label 非空：首个 run 写 label（保留其加粗等格式），正文写入第二个 run
                  （用于「加粗标签：正文」同段的模板）。
    """
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    if label is None:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ''
        return
    runs[0].text = label
    if len(runs) >= 2:
        runs[1].text = text
        for r in runs[2:]:
            r.text = ''
    else:
        new_r = copy.deepcopy(runs[0]._r)
        p._p.append(new_r)
        Run(new_r, p).text = text


def apply_substitutions(p, subs):
    """
    段内全局词替换。subs 为 [(旧, 新), ...]，**长串应排在前面**。
    - 含域代码的段落跳过 run 合并，避免破坏目录。
    - 先按 run 替换；若整段仍残留关键词，再合并 run 后替换。
    """
    if not paragraph_text(p).strip():
        return
    for old, new in subs:
        if old and old in paragraph_text(p):
            for r in p.runs:
                if r.text:
                    r.text = r.text.replace(old, new)
    if has_field_code(p):
        return
    before = paragraph_text(p)
    after = before
    for old, new in subs:
        if old:
            after = after.replace(old, new)
    if after != before and p.runs:
        p.runs[0].text = after
        for r in p.runs[1:]:
            r.text = ''


def insert_paragraphs_after(p, texts, drop_numbering=True):
    """
    在段落 p 之后依次插入新段落；新段深拷贝自 p，
    - drop_numbering=True 时去掉自动编号（numPr），避免正文继承标题编号；
    - 沿用 p 末个 run 的 rPr，保持字号字体。
    """
    prev = p._p
    for txt in texts:
        new_p = copy.deepcopy(p._p)
        pPr = new_p.find(qn('w:pPr'))
        if pPr is not None and drop_numbering:
            numPr = pPr.find(qn('w:numPr'))
            if numPr is not None:
                pPr.remove(numPr)
        runs = new_p.findall(qn('w:r'))
        rPr = None
        if runs:
            last = runs[-1].find(qn('w:rPr'))
            if last is not None:
                rPr = copy.deepcopy(last)
        for r in runs:
            new_p.remove(r)
        for tag in (qn('w:bookmarkStart'), qn('w:bookmarkEnd')):
            for e in new_p.findall(tag):
                new_p.remove(e)
        r = OxmlElement('w:r')
        if rPr is not None:
            r.append(rPr)
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = txt
        r.append(t)
        new_p.append(r)
        prev.addnext(new_p)
        prev = new_p


# --------------------------------------------------------------------------
# 表格填充
# --------------------------------------------------------------------------
def set_cell_text(cell, text):
    """写单元格：只保留一个段落、一个 run，避免残留旧内容"""
    ps = cell.paragraphs
    for extra in ps[1:]:
        extra._p.getparent().remove(extra._p)
    p = ps[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)


def fill_table(doc, idx, header_rows, rows, header_texts=None, row0_texts=None):
    """
    填充第 idx 个表格（0 为封面表，通常不改）：
    - 保留前 header_rows 行表头，其余为数据行；
    - 数据行多于现存 → 深拷贝末行补足；少于 → 删除多余行；
    - header_texts / row0_texts 可选覆盖表头与首行文字。
    """
    t = doc.tables[idx]
    if row0_texts:
        for j, txt in enumerate(row0_texts):
            if j < len(t.rows[0].cells):
                set_cell_text(t.rows[0].cells[j], txt)
    if header_texts:
        hr = t.rows[header_rows - 1]
        for j, txt in enumerate(header_texts):
            if j < len(hr.cells):
                set_cell_text(hr.cells[j], txt)
    have = len(t.rows) - header_rows
    while have < len(rows):
        t._tbl.append(copy.deepcopy(t.rows[-1]._tr))
        have += 1
    while have > len(rows):
        t._tbl.remove(t.rows[-1]._tr)
        have -= 1
    for i, row in enumerate(rows):
        tr = t.rows[header_rows + i]
        for j, val in enumerate(row):
            if j < len(tr.cells):
                set_cell_text(tr.cells[j], str(val))


def clone_table_props(src_table, dst_table, drop_float=True):
    """把模板表的 tblPr 复制给新表（保留边框/单元格边距），可选去掉浮动定位"""
    tp = src_table._tbl.find(qn('w:tblPr'))
    if tp is None:
        return
    new_pr = copy.deepcopy(tp)
    if drop_float:
        f = new_pr.find(qn('w:tblpPr'))
        if f is not None:
            new_pr.remove(f)
    old = dst_table._tbl.find(qn('w:tblPr'))
    if old is not None:
        dst_table._tbl.replace(old, new_pr)
    else:
        dst_table._tbl.insert(0, new_pr)


# --------------------------------------------------------------------------
# 封面 / 分页
# --------------------------------------------------------------------------
def trim_cover_spacers(doc, first_cell_text='文件编号'):
    """删除封面表与首个正文块之间的纯空段；返回删除数"""
    body = doc.element.body
    cover = find_cover_table(body, first_cell_text)
    if cover is None:
        return 0
    started, removed = False, 0
    for el in list(body):
        if el is cover:
            started = True
            continue
        if not started:
            continue
        if el.tag == qn('w:p'):
            if has_section_break_el(el):
                break
            if ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip():
                break
            if el.findall('.//' + qn('w:br')) or el.findall('.//' + qn('w:drawing')):
                break
            body.remove(el)
            removed += 1
        elif el.tag == qn('w:tbl'):
            break
        else:
            continue          # ★ bookmark* 等：跳过而非 break
    return removed


def has_section_break_el(p_el):
    pPr = p_el.find(qn('w:pPr'))
    return pPr is not None and pPr.find(qn('w:sectPr')) is not None


def ensure_cover_page_break(doc, first_cell_text='文件编号'):
    """
    封面表后若无分节符/分页符，则紧随其后插入分页符，保证封面独立成页。
    返回是否插入。
    """
    body = doc.element.body
    cover = find_cover_table(body, first_cell_text)
    if cover is None:
        return False
    started = False
    for el in body:
        if el is cover:
            started = True
            continue
        if not started:
            continue
        if el.tag == qn('w:p'):
            if has_section_break_el(el):
                return False
            if el.findall('.//' + qn('w:br')):
                return False
            if ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip():
                break
        elif el.tag == qn('w:tbl'):
            break
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r = OxmlElement('w:r')
    r.append(br)
    p = OxmlElement('w:p')
    p.append(r)
    cover.addnext(p)
    return True


def clean_body_keeping_section_breaks(doc, from_cover=True,
                                      first_cell_text='文件编号'):
    """
    清空封面之后的正文块，但**保留一个含分节符的段落**与文末 sectPr，
    以免页眉页脚及分节结构丢失（仅在「重建正文」方案下使用）。
    """
    body = doc.element.body
    cover = find_cover_table(body, first_cell_text) if from_cover else None
    started = cover is None
    kept_break = False
    for el in list(body):
        if el is cover:
            started = True
            continue
        if not started:
            continue
        if el.tag == qn('w:sectPr'):
            continue
        if el.tag == qn('w:p') and has_section_break_el(el):
            if not kept_break:
                for child in list(el):
                    if child.tag != qn('w:pPr'):
                        el.remove(child)
                kept_break = True
                continue
        body.remove(el)


def insert_page_break(paragraph_el):
    """在指定段落元素之后插入一个分页符段落"""
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r = OxmlElement('w:r')
    r.append(br)
    p = OxmlElement('w:p')
    p.append(r)
    paragraph_el.addnext(p)


def trim_trailing_empty_paragraphs(doc, keep=1):
    """删除文末（最后一个块/sectPr 之前）的纯空段，避免渲染出末尾空白页。

    模板常在正文结束后遗留若干空段（元良 RD-026 末尾 4 个 → 第 5 页只有页眉）。
    keep=1 表示至少保留一个空段，维持段落标记。返回删除数。
    """
    body = doc.element.body
    # 文末 sectPr 之前的块序列
    blocks = [el for el in body if el.tag != qn('w:sectPr')]
    tail = []
    for el in reversed(blocks):
        if el.tag == qn('w:p') and not ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip() \
                and not el.findall('.//' + qn('w:br')) \
                and not el.findall('.//' + qn('w:drawing')) \
                and not el.findall('.//' + qn('w:pict')) \
                and not has_section_break_el(el):
            tail.append(el)
        else:
            break
    removed = 0
    for el in tail:
        if removed >= max(0, len(tail) - keep):
            break
        body.remove(el)
        removed += 1
    return removed


# --------------------------------------------------------------------------
# 图片：图题驱动的插图替换与校验
# 用法：specs = [(png, '图 1  xxx', anchor_heading, width_in), ...]
#      n = replace_figures(doc, specs)     # 删旧图 + 依序插入
#      drop_orphan_image_rels(doc)         # 清掉包内旧图残留（瘦身）
#      ok, rows = verify_figures(doc, [s[0] for s in specs])
# --------------------------------------------------------------------------
# 图片引用只有两种形态：DrawingML(a:blip@r:embed) 与 VML(v:imagedata@r:id)。
# 模板文档两者常混用（见 pitfalls.md「图片两种形式都要认」）——
# 只认 DrawingML 会漏图，甚至把仍被 VML 引用的图当成孤立件删掉。
# 注意：python-docx 的 nsmap 没有 v 前缀，VML 必须用 Clark 记法（{命名空间}标签）。
VML_NS = 'urn:schemas-microsoft-com:vml'
IMAGE_REF_TAGS = (
    (qn('a:blip'), qn('r:embed')),
    ('{%s}imagedata' % VML_NS, qn('r:id')),
)
TXBX = qn('w:txbxContent')          # 文本框内容容器（字体统一要跳过）
IMAGE_EXTS = ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff', '.emf', '.wmf')


def image_rids(p):
    """段落里全部图片引用的 rId（DrawingML + VML 都认）；文本框内的也算"""
    rids = []
    for tag, attr in IMAGE_REF_TAGS:
        for el in p._p.findall('.//' + tag):
            rid = el.get(attr)
            if rid:
                rids.append(rid)
    return rids


def has_inline_image(p):
    """段落是否含图片（DrawingML 或 VML）"""
    return bool(image_rids(p))


def iter_image_paragraphs(doc):
    """按文档顺序返回含图片的段落（两种图片形式都算）"""
    return [p for p in doc.paragraphs if has_inline_image(p)]


def paragraph_image_sha256(doc, p):
    """取段落内嵌图片的 sha256（图替换后的二进制校验用）"""
    rid = image_rids(p)[0]
    return hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest()


def find_body_heading(doc, text):
    """按前缀在**正文**定位标题段落；跳过目录（toc 样式）段落。

    必须跳过目录：目录条目与正文标题同文本，简单前缀匹配会插错位置。
    """
    for p in doc.paragraphs:
        style = p.style.name if p.style else ''
        if style.lower().startswith('toc'):
            continue
        if paragraph_text(p).strip().startswith(text):
            return p
    return None


def clear_figures(doc, caption_prefix='图 '):
    """删除所有内联图片段落 + 紧随其后的居中图题段落，返回删除条数"""
    doomed = []
    for p in doc.paragraphs:
        if has_inline_image(p):
            doomed.append(p)
        elif paragraph_text(p).strip().startswith(caption_prefix) \
                and p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            doomed.append(p)
    for p in doomed:
        p._p.getparent().remove(p._p)
    return len(doomed)


def insert_figure_before(doc, anchor_para, img_path, caption, width_in,
                         caption_size_pt=10.5, caption_bold=True):
    """在 anchor 段落前插入「居中图片 + 居中图题」两个段落。

    doc.add_picture 只能追加到文末，这里靠 addprevious 把它挪到锚点前。
    """
    doc.add_picture(img_path, width=Inches(width_in))
    pic = doc.paragraphs[-1]
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    anchor_para._p.addprevious(pic._p)

    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        r.font.size = Pt(caption_size_pt)
        r.bold = caption_bold
    anchor_para._p.addprevious(cap._p)


def replace_figures(doc, specs, caption_size_pt=10.5):
    """重建全部插图：先清旧图，再按 specs 依序插到各锚点前。

    specs: [(png_path, caption, anchor_heading_text, width_in), ...]
    同一锚点会被插入多张时，按 specs 顺序依次排好。
    """
    clear_figures(doc)
    anchors = {}
    for _, _, anchor, _ in specs:
        if anchor not in anchors:
            t = find_body_heading(doc, anchor)
            if t is None:
                raise KeyError('未在正文找到锚点标题: %s' % anchor)
            anchors[anchor] = t
    for png, caption, anchor, width_in in specs:
        insert_figure_before(doc, anchors[anchor], png, caption, width_in, caption_size_pt)
    return len(specs)


def drop_orphan_image_rels(doc):
    """删除未被引用的图片关系（瘦身）。

    删段落不会连带删包内图片部件，残留旧图会让文件虚胖（曾 1.5MB→2.4MB）。
    python-docx 保存时只写出「可达」部件，所以断开关系即可，无需改 zip。

    ★ 判定「已引用」必须同时认 DrawingML 与 VML：只认 a:blip 会把 VML
      (v:imagedata) 正在用的图当成孤立件删掉，保存后图片丢失、引用悬空。
    """
    used = set()
    for p in doc.paragraphs:
        for rid in image_rids(p):
            rel = doc.part.rels.get(rid)
            if rel is not None:
                used.add(os.path.basename(rel.target_ref or ''))
    dropped = []
    for rid, rel in list(doc.part.rels.items()):
        base = os.path.basename(rel.target_ref or '')
        if base.startswith('image') and base.lower().endswith(IMAGE_EXTS) and base not in used:
            doc.part.drop_rel(rid)
            dropped.append(base)
    return dropped


def png_size(path):
    """读 PNG 头取 (宽, 高) 像素"""
    with open(path, 'rb') as f:
        head = f.read(24)
    return struct.unpack('>II', head[16:24])


def verify_figures(doc, expected_pngs):
    """校验图组：数量/顺序 + 二进制哈希 + 显示宽高比（防拉伸）。

    返回 (ok, rows)，rows 每项为 dict(png, sha_ok, ratio_ok, disp_w, disp_h, ratio)。
    """
    rows = []
    ok = True
    for p, png in zip(iter_image_paragraphs(doc), expected_pngs):
        blob = doc.part.related_parts[image_rids(p)[0]].blob
        sha_ok = hashlib.sha256(blob).hexdigest() == \
            hashlib.sha256(open(png, 'rb').read()).hexdigest()
        # VML 图没有 wp:inline/extent，无法核显示比例，只核二进制（避免误报崩溃）
        inl = p._p.findall('.//' + qn('wp:inline'))
        ext = inl[0].find(qn('wp:extent')) if inl else None
        if ext is None:
            ok = ok and sha_ok
            rows.append(dict(png=os.path.basename(png), sha_ok=sha_ok, ratio_ok=True,
                             disp_w=None, disp_h=None, ratio=None))
            continue
        cx, cy = int(ext.get('cx')), int(ext.get('cy'))
        pw, ph = png_size(png)
        ratio = cx / cy
        ratio_ok = abs(pw / ph - ratio) < 0.01          # 容差 1%
        ok = ok and sha_ok and ratio_ok
        rows.append(dict(png=os.path.basename(png), sha_ok=sha_ok, ratio_ok=ratio_ok,
                         disp_w=cx / 914400.0, disp_h=cy / 914400.0, ratio=ratio))
    if len(rows) != len(expected_pngs):
        ok = False
    return ok, rows


def assert_writable(path):
    """确认文件未被 Word/WPS 独占锁定；被占用时抛 IOError。

    python-docx 保存到被打开的 .docx 会报 PermissionError，且此时改动已丢，
    所以动笔前先探一次。
    """
    try:
        with open(path, 'rb+'):
            pass
    except OSError as e:
        raise IOError('目标文档被其它进程占用，请先关闭：%s' % path) from e


# --------------------------------------------------------------------------
# 段落缩进规范化 / 旧式图形清理
#   模板中常见「畸形首行缩进」：firstLine 远大于 2 字符（如 2040/3420 twip），
#   并常与负的 leftChars 同时出现 → 渲染成「首行缩进一大截、折行回到左边界」。
#   该缺陷在 06 项目模板、血糖仪 058/062、元良三个项目上均复现，故独立成原语。
#   注意：**不要把正文文字写进畸形段落**；应留空该段落，文字写入新建的正常段落。
# --------------------------------------------------------------------------
def normalize_paragraph_indent(doc, max_first_line=1440, first_line_chars=200,
                               first_line=480, drop_zero=False):
    """把畸形首行缩进规范化。返回 (修改数, 清理空段数)。

    - 段落有文字：仅替换 w:ind 为 (firstLineChars, firstLine)，保留其余格式；
    - 段落为空  ：drop_zero=True 时直接删除该空段，否则同样规范化。
    """
    from docx.oxml import OxmlElement
    fixed = dropped = 0
    for p in list(doc.paragraphs):
        pPr = p._p.find(qn('w:pPr'))
        if pPr is None:
            continue
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            continue
        fl = ind.get(qn('w:firstLine'))
        lc = ind.get(qn('w:leftChars'))
        bad = (fl and abs(int(fl)) > max_first_line) or (lc and int(lc) < 0)
        if not bad:
            continue
        empty = not ''.join(t.text or '' for t in p._p.iter(qn('w:t'))).strip()
        if empty and drop_zero:
            p._p.getparent().remove(p._p)
            dropped += 1
            continue
        pPr.remove(ind)
        new = OxmlElement('w:ind')
        new.set(qn('w:firstLineChars'), str(first_line_chars))
        new.set(qn('w:firstLine'), str(first_line))
        # 按 schema，w:ind 须排在 w:rPr 之前
        rPr = pPr.find(qn('w:rPr'))
        if rPr is not None:
            rPr.addprevious(new)
        else:
            pPr.append(new)
        fixed += 1
    return fixed, dropped


def clear_legacy_shapes(doc, caption_prefixes=('图', '表'), drop_empty=True):
    """删除旧式 VML 图形段落（w:pict）及其紧随的图题段落。

    模板里继承的图常是 VML（不是 a:blip），`clear_figures` 抓不到；图题
    也可能是「图1 xxx」（无空格），不匹配 `图 ` 前缀。故本函数：
      1) 删掉所有含 w:pict 的段落；
      2) 删掉文本以 caption_prefixes 开头、且紧跟其后的图题段落。
    返回删除段数。
    """
    targets = []
    for i, p in enumerate(doc.paragraphs):
        if p._p.findall('.//' + qn('w:pict')):
            targets.append(p)
            continue
        t = paragraph_text(p).strip()
        if t and any(t.startswith(x) for x in caption_prefixes) and len(t) <= 40:
            prev = doc.paragraphs[i - 1] if i > 0 else None
            if prev is not None and (prev._p.findall('.//' + qn('w:pict'))
                                    or has_inline_image(prev)):
                targets.append(p)
    for p in targets:
        p._p.getparent().remove(p._p)
    return len(targets)


# --------------------------------------------------------------------------
# 文档级读写
# --------------------------------------------------------------------------
def open_template(path):
    return Document(path)


def save_doc(doc, out_dir, filename):
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, filename)
    doc.save(out)
    return out


def extract_text(path):
    """抽取正文 + 表格文本（用于一致性检查）；.doc 请先转换"""
    d = Document(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                if c.text.strip():
                    lines.append(c.text)
    return '\n'.join(lines)


def convert_doc_to_docx(src, dst, prefer=('Word.Application', 'KWPS.Application')):
    """用 Word/WPS COM 把 .doc 转为 .docx；两者都失败则抛异常"""
    import win32com.client as w
    last = None
    for prog in prefer:
        try:
            app = w.Dispatch(prog)
            app.Visible = False
            app.DisplayAlerts = 0
            d = app.Documents.Open(src, ReadOnly=True)
            d.SaveAs2(dst, FileFormat=16)
            d.Close(False)
            app.Quit()
            return dst
        except Exception as e:      # noqa: BLE001
            last = e
            continue
    raise RuntimeError('无法转换 .doc：%s' % last)


# ─────────────────────────────────────────────────────────────────────────────
# 域（Field）/ 目录（TOC）修复与分页一致性（详见 docx-fields-and-toc-repair.md、
# cover-pagination-and-sections.md 第 5/6 节）。四个函数互相独立，可单独调用。
# ─────────────────────────────────────────────────────────────────────────────

def toc_diagnose(doc):
    """诊断目录域健康度。
    返回 keys：
      balance        —— begin-end 差值，>0 表示域未闭合（如 RD-068 缺 end）
      toc_instr      —— 是否找到 TOC 指令（instrText）
      toc_has_begin  —— TOC 指令前是否有未被闭合的 begin（False → 指令裸露，如 RD-026）
      toc_has_separate —— TOC 指令后是否有 separate（False → 结果区未开始，指令会裸露）
      text_has_toc   —— 指令是否被误写成 <w:t> 普通文本
      toc_paras      —— 目录样式段落索引
      entry_has_field—— 目录条目是否含域（False → 纯文本目录，无超链接，如 RD-032）
    """
    from collections import Counter
    body = doc.element.body
    seq = []
    for el in body.iter():
        if el.tag == qn('w:fldChar'):
            seq.append(('fld', el.get(qn('w:fldCharType'))))
        elif el.tag == qn('w:instrText'):
            seq.append(('instr', el.text or ''))
        elif el.tag == qn('w:t'):
            seq.append(('t', el.text or ''))
    c = Counter(v for k, v in seq if k == 'fld')
    pos = [i for i, (k, v) in enumerate(seq) if k == 'instr' and 'TOC' in v]
    toc_instr = bool(pos)
    has_begin = has_sep = False
    if pos:
        i = pos[0]
        depth = 0
        for j in range(i - 1, -1, -1):          # 向前找未被闭合的 begin
            k, v = seq[j]
            if k == 'fld' and v == 'end':
                depth += 1
            elif k == 'fld' and v == 'begin':
                if depth == 0:
                    has_begin = True
                    break
                depth -= 1
        for j in range(i + 1, len(seq)):        # 向后找 separate（到首个 end 为止）
            k, v = seq[j]
            if k == 'fld' and v == 'separate':
                has_sep = True
                break
            if k == 'fld' and v == 'end':
                break
    tocs = [i for i, p in enumerate(doc.paragraphs)
            if (p.style.name or '').lower().startswith('toc')]
    return {
        'begin': c.get('begin', 0), 'separate': c.get('separate', 0), 'end': c.get('end', 0),
        'balance': c.get('begin', 0) - c.get('end', 0),
        'toc_instr': toc_instr,
        'toc_has_begin': has_begin,             # False → 指令裸露
        'toc_has_separate': has_sep,            # False → 结果区未开始
        'text_has_toc': any('TOC' in v for k, v in seq if k == 't'),
        'toc_paras': tocs,
        'entry_has_field': any(doc.paragraphs[i]._p.findall('.//' + qn('w:fldChar')) for i in tocs),
        'applicable': bool(pos) or bool(tocs),        # 本文档是否有目录（无目录 → healthy 为 None）
        'healthy': (None if not (pos or tocs) else
                    bool(pos and has_begin and has_sep and c.get('begin', 0) == c.get('end', 0))),
    }


def _fld_run(fld_type):
    r = OxmlElement('w:r')
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), fld_type)
    r.append(fc)
    return r


def close_toc_field(doc):
    """外层 TOC 域缺 end 时，在目录末条目段落末尾补足；返回补的个数"""
    need = toc_diagnose(doc)['balance']
    if need <= 0:
        return 0
    tocs = toc_diagnose(doc)['toc_paras']
    if not tocs:
        return 0
    last = doc.paragraphs[max(tocs)]
    for _ in range(need):
        last._p.append(_fld_run('end'))
    return need


def fix_stray_toc_instruction(doc):
    """把误写成普通文本（w:t）的 TOC 指令转成真正的域：begin + instrText + separate"""
    for p in doc.paragraphs:
        for r in p._p.findall(qn('w:r')):
            t = r.find(qn('w:t'))
            if t is None or 'TOC' not in (t.text or ''):
                continue
            instr = t.text
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = instr
            r.addprevious(_fld_run('begin'))
            r.addprevious(it)
            r.addprevious(_fld_run('separate'))
            t.text = ''          # 指令不再作为可见文本
            return instr
    return None


def normalize_toc_page_break(doc, pad_min=5):
    """目录末条目→正文首标题之间的垫片空段，改为 1 个显式分页符。返回(删除段数, 是否插入)"""
    ps = doc.paragraphs
    tocs = [i for i, p in enumerate(ps) if (p.style.name or '').lower().startswith('toc')]
    if not tocs:
        return 0, False
    last = max(tocs)
    first_body = None
    for i, p in enumerate(ps):
        if i <= last:
            continue
        pPr = p._p.find(qn('w:pPr'))
        if pPr is not None and pPr.find(qn('w:numPr')) is not None and paragraph_text(p).strip():
            first_body = i
            break
    if first_body is None:
        return 0, False
    empt = [p for p in ps[last + 1:first_body]
            if not paragraph_text(p).strip()
            and not p._p.findall('.//' + qn('w:br'))
            and not p._p.findall('.//' + qn('w:drawing'))
            and not p._p.findall('.//' + qn('w:pict'))
            and not p._p.findall('.//' + qn('w:fldChar'))        # ★ 只含域标记的段不能当空段删
            and not p._p.findall('.//' + qn('w:instrText'))]
    if len(empt) < pad_min:
        return 0, False
    keep = empt[0]
    for r in keep._p.findall(qn('w:r')):
        keep._p.remove(r)
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    keep._p.append(r)
    n = 0
    for p in empt[1:]:
        p._p.getparent().remove(p._p)
        n += 1
    return n, True


def page_size_report(doc):
    """逐节报告页面尺寸/页边距，用于发现"某节被设成横向"导致的页宽不一致"""
    out = []
    for i, s in enumerate(doc.sections):
        e = s._sectPr
        sz = e.find(qn('w:pgSz'))
        mar = e.find(qn('w:pgMar'))
        out.append({
            'section': i + 1,
            'w': sz.get(qn('w:w')) if sz is not None else None,
            'h': sz.get(qn('w:h')) if sz is not None else None,
            'orient': sz.get(qn('w:orient')) if sz is not None else None,
            'margins': tuple(mar.get(qn('w:' + k)) if mar is not None else None
                             for k in ('left', 'right', 'top', 'bottom')),
        })
    return out


def unify_sections(doc, base=0):
    """把各节的 pgSz/pgMar 统一为第 base 节（默认封面节）的值；返回修改记录"""
    b = doc.sections[base]._sectPr
    b_sz, b_mar = b.find(qn('w:pgSz')), b.find(qn('w:pgMar'))
    fixed = []
    for i, s in enumerate(doc.sections):
        if i == base:
            continue
        e = s._sectPr
        sz = e.find(qn('w:pgSz'))
        if sz is None:
            sz = OxmlElement('w:pgSz')
            e.insert(0, sz)
        for k in ('w', 'h', 'orient'):
            bv = b_sz.get(qn('w:' + k)) if b_sz is not None else None
            if bv is None:
                if sz.get(qn('w:' + k)) is not None:
                    del sz.attrib[qn('w:' + k)]
            elif sz.get(qn('w:' + k)) != bv:
                sz.set(qn('w:' + k), bv)
                fixed.append('节%d pgSz.%s→%s' % (i + 1, k, bv))
        mar = e.find(qn('w:pgMar'))
        if mar is not None and b_mar is not None:
            for k in ('left', 'right', 'top', 'bottom'):
                bv = b_mar.get(qn('w:' + k))
                if bv is not None and mar.get(qn('w:' + k)) != bv:
                    mar.set(qn('w:' + k), bv)
                    fixed.append('节%d pgMar.%s→%s' % (i + 1, k, bv))
    return fixed


def clamp_images_to_text_width(doc, max_emu=None):
    """把所有内嵌图缩放到不超过正文宽度（保持纵横比）；返回缩放张数"""
    sec = doc.sections[0]
    if max_emu is None:
        max_emu = sec.page_width - sec.left_margin - sec.right_margin
    n = 0
    for sh in doc.inline_shapes:
        w, h = sh.width, sh.height
        if w and w > max_emu:
            ratio = max_emu / w
            sh.width = Emu(int(max_emu))
            sh.height = Emu(int(h * ratio))
            n += 1
    return n


def refresh_toc_via_word(path, save=True):
    """用 Word 更新目录域/字段（等价于在 Word 里 F9 /『更新目录』）。
    只能在装有 Word/WPS 的环境执行；不刷新则目录页码停留在旧值。"""
    import win32com.client as w
    app = w.Dispatch('Word.Application')
    app.Visible = False
    app.DisplayAlerts = 0
    try:
        d = app.Documents.Open(path)
        for i in range(1, d.TablesOfContents.Count + 1):
            d.TablesOfContents(i).Update()
        d.Fields.Update()
        if save:
            d.Save()
        d.Close(False)
    finally:
        app.Quit()


def export_pdf_via_word(path, pdf_path):
    """LibreOffice 卡死时的兜底：用 Word 只读导出 PDF（得到的即 Word 自身分页）"""
    import win32com.client as w
    app = w.Dispatch('Word.Application')
    app.Visible = False
    app.DisplayAlerts = 0
    try:
        d = app.Documents.Open(path, ReadOnly=True)
        pages = d.ComputeStatistics(2)          # 2 = wdStatisticPages
        d.ExportAsFixedFormat(pdf_path, 17)     # 17 = PDF
        d.Close(False)
        return pages
    finally:
        app.Quit()

# --------------------------------------------------------------------------
# 字体与字号统一（文档族级版式收口；详见 font-and-style-normalization.md）
#   三条实测必要性：
#     ① 必须删掉 rFonts 的 *Theme 属性，否则主题字体优先级高于显式 rFonts，改了不生效；
#     ② 必须单独遍历表格——doc.paragraphs **不含表格内段落**，漏了表格等于没统一；
#     ③ 必须同时改 styles.xml 的 Normal/Heading/Caption，否则用户在 Word 里新打字又回旧字体。
# --------------------------------------------------------------------------
THEME_FONT_ATTRS = ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'ctheme', 'cstheme')


def apply_rpr_font(rPr, east_asia=None, latin=None, sz=None):
    """在给定 `w:rPr` 上**就地**写字体槽与字号。

    ★ 不要为了复用 run 级函数而把 rPr `append` 到临时 `w:r` 上：lxml 的 append 是
      **移动**语义，会把 rPr 从样式元素里搬走，结果 Heading 的字号定义整段丢失，
      而且不报错、只在渲染时表现为"章标题变成正文大小"。
    """
    if east_asia or latin:
        rf = rPr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rPr.insert(0, rf)
        if latin:
            for at in ('ascii', 'hAnsi', 'cs'):
                rf.set(qn('w:' + at), latin)
        if east_asia:
            rf.set(qn('w:eastAsia'), east_asia)
        for at in THEME_FONT_ATTRS:
            if rf.get(qn('w:' + at)):
                del rf.attrib[qn('w:' + at)]
    if sz:
        for tag in ('w:sz', 'w:szCs'):
            el = rPr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                rPr.append(el)
            el.set(qn('w:val'), str(sz))
    return rPr


def style_font(doc, name, east_asia=None, latin=None, sz=None):
    """改样式默认值（Normal / Heading 1~4 / Caption / Body Text）。样式名不存在时返回 False。"""
    try:
        el = doc.styles[name]._element
    except KeyError:
        return False
    rPr = el.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        el.append(rPr)
    apply_rpr_font(rPr, east_asia, latin, sz)
    return True


def _in_textbox(el):
    anc = el.getparent()
    while anc is not None:
        if anc.tag == TXBX:
            return True
        anc = anc.getparent()
    return False


def iter_zone_paragraphs(doc, cover_first_cell='文件编号'):
    """按版式分区产出 (paragraph, zone)：
    cover / head / caption / table / body / hf（页眉页脚）。

    分区是字号统一的前提——标题、封面、页眉**只该改字体不该改字号**，
    否则章标题会和正文一样大、封面版式塌陷。
    """
    sect_i = next((i for i, p in enumerate(doc.paragraphs) if has_section_break(p)), 0)
    cover = find_cover_table(doc.element.body, cover_first_cell)
    for i, p in enumerate(doc.paragraphs):
        pPr = p._p.find(qn('w:pPr'))
        style = (p.style.name or '').lower()
        t = ''.join(x.text or '' for x in p._p.iter(qn('w:t')))
        t = re.sub(r'\s+', '', t or '')
        if i < sect_i:
            zone = 'cover'
        elif style.startswith('heading') or style.startswith('toc')                 or style.startswith('caption')                 or (pPr is not None and pPr.find(qn('w:numPr')) is not None):
            zone = 'head'
        elif t[:1] in ('图', '表') and len(t) <= 40:
            zone = 'caption'
        else:
            zone = 'body'
        yield p, zone
    for tb in doc.tables:
        z = 'cover' if tb._tbl is cover else 'table'
        for row in tb.rows:
            seen = set()
            for c in row.cells:
                if c._tc in seen:
                    continue
                seen.add(c._tc)
                for p in c.paragraphs:
                    yield p, z
    for s in doc.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                   s.even_page_header, s.even_page_footer):
            if hf is not None:
                for p in hf.paragraphs:
                    yield p, 'hf'


def unify_document_fonts(doc, east_asia='宋体', latin='Times New Roman',
                         body_sz=None, table_sz=None, caption_sz=None,
                         styles=('Normal', 'Heading 1', 'Heading 2', 'Heading 3',
                                 'Heading 4', 'Caption', 'Body Text'),
                         cover_first_cell='文件编号'):
    """统一整份文档字体，并按分区设字号（字号单位=半磅，None 表示该分区不动）。

    返回 (改字体的 run 数, 设字号的 run 数, 改的样式数)。文本框内的 run 一律跳过：
    那是表单/示意图的排版，改字号会撑破框。
    """
    sizes = {'body': body_sz, 'table': table_sz, 'caption': caption_sz}
    n_font = n_sz = 0
    for p, zone in iter_zone_paragraphs(doc, cover_first_cell):
        sz = sizes.get(zone)
        for r in p._p.findall(qn('w:r')):
            if _in_textbox(r):
                continue
            rPr = r.find(qn('w:rPr'))
            if rPr is None:
                rPr = OxmlElement('w:rPr')
                r.insert(0, rPr)
            apply_rpr_font(rPr, east_asia, latin, sz)
            n_font += 1
            if sz:
                n_sz += 1
    n_style = sum(1 for nm in styles
                  if style_font(doc, nm, east_asia, latin,
                                body_sz if nm == 'Normal' else None))
    return n_font, n_sz, n_style


# --------------------------------------------------------------------------
# 表格重建原语（fill_table 之外的两种结构改动）
# --------------------------------------------------------------------------
def table_grid_cols(tbl):
    """现有列数：**必须看 tblGrid**，不能用末行的 tc 数——合并单元格会让 tc 数少于列数，
    据此判断会漏调列，产出「表头 5 列、数据行 4 格」的错位表。"""
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        return len(grid.findall(qn('w:gridCol')))
    trs = tbl.findall(qn('w:tr'))
    return len(trs[-1].findall(qn('w:tc'))) if trs else 0


def flatten_table_rows(tbl, keep_head=0):
    """去掉纵向合并与数据行的横向合并。

    克隆末行做新数据行时，若该带 `vMerge`，首个新行没有可继承的上方 tc，
    python-docx 读 `row.cells` 会直接抛
    `ValueError: no `tc` element at grid_offset=…`（不是内容错误，是遍历失败）。
    """
    for ri, tr in enumerate(tbl.findall(qn('w:tr'))):
        for tc in tr.findall(qn('w:tc')):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                continue
            for vm in tcPr.findall(qn('w:vMerge')):
                tcPr.remove(vm)
            if ri >= keep_head:
                for gs in tcPr.findall(qn('w:gridSpan')):
                    tcPr.remove(gs)


def substitute_wt_nodes(p, subs):
    """段落级词替换的**图形安全**版本。

    `docx_ops.apply_substitutions` 在整段仍残留关键词时会走"合并 run"分支
    （`p.runs[0].text = after`），而 `run.text` 赋值会清空该 run 的其它子节点 ——
    含 `w:pict`/`w:drawing` 的段落（表单文本框、示意图）会被连带抹掉，且不报错。
    含图形的段落只用本函数：逐 `w:t` 节点替换，不合并 run。
    """
    graphic = bool(p._p.findall('.//' + qn('w:pict'))
                   or p._p.findall('.//' + qn('w:drawing')))
    if not graphic:
        apply_substitutions(p, subs)
        return False
    for t in p._p.iter(qn('w:t')):
        if not t.text:
            continue
        new = t.text
        for old, rep in subs:
            if old and old in new:
                new = new.replace(old, rep)
        if new != t.text:
            t.text = new
            t.set(qn('xml:space'), 'preserve')
    return True


def release_orphan_word_locks(match=(), verbose=False):
    """释放"被杀掉的导出脚本"留下的无头 Word 锁。

    场景：用 Word COM 导 PDF 的脚本被中途终止 → finally 不执行 → WINWORD 仍开着产物文件，
    下一轮生成器写同名文件直接 PermissionError。判据要保守：**只处理 Visible=False 的实例**
    （用户自己开的 Word 一定是可见的），且只关闭文件名命中 match 前缀的文档，
    用 SaveChanges=0 不写盘（Word 在导出期会给文档打临时改动，写回去等于污染产物）。
    返回 (关闭的文档数, 是否退出实例)。
    """
    import win32com.client as w
    try:
        app = w.GetActiveObject('Word.Application')
    except Exception:
        return 0, False
    if getattr(app, 'Visible', True):
        return 0, False                       # 用户正在用：不动
    closed = 0
    for d in list(app.Documents):
        nm = d.Name
        if match and not any(str(nm).startswith(m) for m in match):
            continue
        d.Close(SaveChanges=0)
        closed += 1
        if verbose:
            print('closed', nm)
    quit_ok = False
    try:
        if app.Documents.Count == 0:
            app.Quit()
            quit_ok = True
    except Exception:
        pass
    return closed, quit_ok


# --------------------------------------------------------------------------
# 版式闸门：分页 / 标题树 / 图文同页 / 渲染判据
# --------------------------------------------------------------------------

_NUM_HEADING_RE = re.compile(
    r'^(第\s*[0-9一二三四五六七八九十百]+\s*[章节篇部分]'      # 第X章 / 第X节
    r'|\d+(?:\.\d+){0,2}\s'                                    # 1 / 1.1 / 1.1.1
    r'|\d+(?:\.\d+){1,2}(?=[^\d.\-]))')                        # 1.1名称；排除 1.1-（图号）
_CAPTION_RE = re.compile(r'^([图表]\s*[\d\-.]|\d+(?:\.\d+)*-\d+)')  # 图1 / 表2-1 / 2.7-1 流程图


def _style_outline_map(doc):
    """{styleId: (显示名, 样式自带 outlineLvl 或 None, basedOn)}。"""
    m = {}
    for s in doc.styles.element.findall(qn('w:style')):
        nm = s.find(qn('w:name'))
        ol = s.find('.//' + qn('w:outlineLvl'))
        bo = s.find(qn('w:basedOn'))
        m[s.get(qn('w:styleId'))] = (
            nm.get(qn('w:val')) if nm is not None else '?',
            int(ol.get(qn('w:val'))) if ol is not None else None,
            bo.get(qn('w:val')) if bo is not None else None)
    return m


def effective_outline(p_el, smap, _depth=0):
    """段落生效大纲级别：直接 pPr 优先，否则沿样式链；None = 正文级。

    必须沿链判定（§34）：值常在样式里而不在段落自己手上。
    """
    pPr = p_el.find(qn('w:pPr'))
    if pPr is None:
        return None
    direct = pPr.find(qn('w:outlineLvl'))
    if direct is not None:
        v = int(direct.get(qn('w:val')))
        return None if v > 8 else v
    st = pPr.find(qn('w:pStyle'))
    sid = st.get(qn('w:val')) if st is not None else None
    while sid and sid in smap and _depth <= 8:
        name, ol, based = smap[sid]
        if ol is not None:
            return None if ol > 8 else ol
        sid = based
        _depth += 1
    return None


def pagination_audit(doc, cover_texts=()):
    """分页 / 标题树体检（只读）。返回 dict，各项为空即通过。

    对应 pitfalls §23 §24 §25 §38 §39 §41。`cover_texts` 传封面区文档名
    （归一化去空白后比较），用于识别"封面短行被写成一级标题"。
    """
    smap = _style_outline_map(doc)
    body = doc.element.body
    kids = list(body.iterchildren())
    # 以下三个 helper 作用于 CT_P 元素（库内 paragraph_text/has_graphic 收 Paragraph 对象）
    etxt = lambda el: ''.join(t.text or '' for t in el.iter(qn('w:t')))
    egra = lambda el: bool(el.findall('.//' + qn('w:drawing'))
                           or el.findall('.//' + qn('w:pict')))

    def flag(el, tag):
        """布尔型 pPr 子元素：必须判 w:val，写成 val="0"/"false" 是**关闭**。"""
        if el is None:
            return False
        e = el.find(qn(tag))
        if e is None:
            return False
        v = (e.get(qn('w:val')) or '1').strip().lower()
        return v not in ('0', 'false', 'off', 'none')

    def in_field(el):
        """段落是否落在某个域内（目录/索引条目自带编号文本与页码，§41 必须先排除）。"""
        if el.findall('.//' + qn('w:fldChar')) or el.findall('.//' + qn('w:instrText')):
            return True
        st = el.find('.//' + qn('w:pStyle'))
        sid = st.get(qn('w:val')) if st is not None else ''
        nm = smap.get(sid, ('', None, None))[0]
        return nm.startswith('TOC') or nm.startswith('目录')

    out = dict(hard_page_breaks=0, stacked_breaks=[], body_page_break=[],
               keep_on_body=[], chapter_pad=[], outline_anomaly=[],
               numbered_not_in_tree=[], toc_field=False)

    # 封面区边界：封面信息表之前的空段属封面留白，不是"章前顶页回车"（§23 只管正文区）
    cover_end = len(kids)
    for k, c in enumerate(kids):
        if c.tag.split('}')[-1] == 'tbl' and '文件编号' in etxt(c):
            cover_end = k
            break

    for b in doc.element.iter(qn('w:br')):
        if b.get(qn('w:type')) == 'page':
            out['hard_page_breaks'] += 1
    out['toc_field'] = bool(re.search(r'<w:instrText[^>]*>\s*TOC',
                                      doc.element.xml))

    for i, c in enumerate(kids):
        if c.tag.split('}')[-1] != 'p':
            continue
        pPr = c.find(qn('w:pPr'))
        t = etxt(c).strip()
        nt = re.sub(r'\s+', '', t)
        lvl = effective_outline(c, smap)
        pbb = flag(pPr, 'w:pageBreakBefore')
        hard = any(b.get(qn('w:type')) == 'page'
                   for b in c.iter(qn('w:br')))
        if hard and pbb:
            out['stacked_breaks'].append(t[:30] or '<空段>')          # §39
        if pbb and lvl is None and t:
            out['body_page_break'].append(t[:30])                     # §39
        if lvl is None and pPr is not None and (
                flag(pPr, 'w:keepNext') or flag(pPr, 'w:keepLines')):
            out['keep_on_body'].append(                               # §38
                ('<图/空段>' if egra(c) else '') + t[:30])
        if lvl == 0:
            n = 0
            j = i - 1
            while j >= 0 and kids[j].tag.split('}')[-1] == 'p' \
                    and not etxt(kids[j]).strip() and not egra(kids[j]):
                # 遇分节符载体段即停：其前空段属上一节，不是本章的顶页回车（§23 边界1）
                if kids[j].find('.//' + qn('w:sectPr')) is not None:
                    n = 0
                    break
                n += 1
                j -= 1
            if n >= 2 and i >= cover_end:
                out['chapter_pad'].append((t[:24], n))                # §23
            if nt in {re.sub(r'\s+', '', x) for x in cover_texts}:
                out['outline_anomaly'].append(('封面短行带一级大纲', t[:24]))
        if lvl is not None and not t and not in_field(c):
            out['outline_anomaly'].append(('空标题段', ''))            # §24
        if lvl is None and t and not in_field(c) \
                and _NUM_HEADING_RE.match(t) and not _CAPTION_RE.match(t):
            out['numbered_not_in_tree'].append(t[:30])                # §41
    return out


def page_ink_coverage(pdf_path, dpi=50, band=0.08, cutoff=200):
    """逐页返回"正文区墨占比(%)"，用于空白页判据（§40）。

    中文 PDF 常为子集字体且无 ToUnicode，pdftotext 只吐数字/拉丁，
    "字数"判据会把正常纯中文页误判为空白页——墨占比不受字体影响。
    经验档位：封面 2~3%，正文 3~13%，真空白页 0.0%。
    """
    import subprocess, tempfile, os, shutil
    tmp = tempfile.mkdtemp()
    try:
        subprocess.run(['pdftoppm', '-png', '-r', str(dpi), pdf_path,
                        os.path.join(tmp, 'p')], check=True, capture_output=True)
        pages = sorted(f for f in os.listdir(tmp) if f.endswith('.png'))
        res = []
        try:
            from PIL import Image
        except ImportError:                                     # 退化：PyMuPDF
            import fitz
            doc = fitz.open(pdf_path)
            for pg in doc:
                pix = pg.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
                w, h = pix.width, pix.height
                buf = pix.samples
                y0, y1 = int(h * band), int(h * (1 - band))
                seg = buf[y0 * w:(y1) * w]
                res.append(round(sum(1 for v in seg if v < cutoff) / len(seg) * 100, 3))
            doc.close()
            return res
        for f in pages:
            im = Image.open(os.path.join(tmp, f)).convert('L')
            w, h = im.size
            crop = im.crop((0, int(h * band), w, int(h * (1 - band))))
            px = list(crop.getdata())
            res.append(round(sum(1 for v in px if v < cutoff) / len(px) * 100, 3))
        return res
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def grade_pages_by_ink(coverage, blank=0.15, thin=0.8):
    """按 §37 分级 + §40 墨占比判据给出 (阻断项, 提示项)。"""
    block, warn = [], []
    n = len(coverage)
    for i, c in enumerate(coverage, 1):
        if c < blank:
            block.append('第%d页正文区无墨（空白页）' % i)
        elif i < n and c < thin:
            block.append('第%d页仅%.2f%%墨（疑似顶页/分页错位）' % (i, c))
        elif i == n and c < thin:
            warn.append('第%d页%.2f%%墨（末页溢出，交人决定）' % (i, c))
    return block, warn
